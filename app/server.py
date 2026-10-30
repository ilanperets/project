import json, os, re, shutil, time, uuid
from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from . import classify as C, engines as E, boq as B, brain as BR, derive as DV, symbols as SY, mapper as MP
import datetime

ROOT=os.environ.get('KAMUYOT_DATA',os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),'data'))
PROJ=os.path.join(ROOT,'projects'); os.makedirs(PROJ,exist_ok=True)
RULES=os.path.join(ROOT,'rules.json')
MEM=os.path.join(ROOT,'prices_memory.json')
SYMF=os.path.join(ROOT,'symbols.json'); LMAPF=os.path.join(ROOT,'layermap.json')
def _jl(p,d):
    try: return json.load(open(p,encoding='utf-8'))
    except Exception: return d
def _js(p,v): os.makedirs(ROOT,exist_ok=True); json.dump(v,open(p,'w',encoding='utf-8'),ensure_ascii=False)
app=FastAPI(title='מערכת כתבי כמויות')

def _rules():
    try: return json.load(open(RULES,encoding='utf-8'))
    except Exception: return []
def _pdir(pid):
    if not re.fullmatch(r'[a-f0-9]{8}',pid): raise HTTPException(404,'פרויקט לא נמצא')
    return os.path.join(PROJ,pid)
def _load(pid):
    p=os.path.join(_pdir(pid),'project.json')
    if not os.path.exists(p): raise HTTPException(404,'פרויקט לא נמצא')
    pr=json.load(open(p,encoding='utf-8'))
    for k,v in (('materials',[]),('prices',{}),('mprices',{}),('params',{}),('bills',[]),('orders',[]),('quotes',[]),('needs',[]),('log',[]),('events',[]),('pending',[]),('stale',False),('checks',{}),('version',1),('status','received'),('client',''),('measurer',''),('reviewer',''),('due',''),('created',''),('updated',''),('changes',None),('delivered_at','')): pr.setdefault(k,v)
    return pr
def _ev(pr,t): pr.setdefault('events',[]).append(dict(t=t,at=datetime.datetime.now().isoformat(timespec='minutes'))); pr['events']=pr['events'][-40:]
def _save(pr):
    pr['updated']=datetime.datetime.now().isoformat(timespec='minutes')
    json.dump(pr,open(os.path.join(_pdir(pr['id']),'project.json'),'w',encoding='utf-8'),ensure_ascii=False)

@app.get('/api/health')
def health(): return dict(ok=True,types=C.TYPES)

@app.get('/api/projects')
def projects():
    out=[]
    for d in sorted(os.listdir(PROJ)):
        try: pr=_load(d); out.append(dict(id=pr['id'],name=pr['name'],sheets=len(pr['sheets']),has_boq=bool(pr.get('boq')),status=pr.get('status')))
        except Exception: pass
    return out

@app.post('/api/projects')
def create(body:dict):
    pid=uuid.uuid4().hex[:8]; os.makedirs(os.path.join(_pdir(pid),'files'))
    now=datetime.datetime.now().isoformat(timespec='minutes'); pr=dict(id=pid,name=(body.get('name') or 'פרויקט חדש').strip()[:80],client=(body.get('client') or '').strip()[:60],measurer=(body.get('measurer') or '').strip()[:40],reviewer=(body.get('reviewer') or '').strip()[:40],due=(body.get('due') or '')[:10],status='received',version=1,created=now,events=[dict(t='הפרויקט נפתח',at=now)],checks={},scale=None,sheets=[],boq=None,needs=[],log=[],materials=[],prices={},mprices={},params={},bills=[],orders=[],quotes=[]); _save(pr); return pr

@app.post('/api/reset')
def reset(body:dict):
    """wipe all projects; optionally also the learned price memory and classification rules"""
    shutil.rmtree(PROJ,ignore_errors=True); os.makedirs(PROJ,exist_ok=True)
    if body.get('memory'):
        try: os.remove(MEM)
        except OSError: pass
    if body.get('rules'):
        for f in (RULES,SYMF,LMAPF):
            try: os.remove(f)
            except OSError: pass
    return dict(ok=True)

STATUSES=['received','measuring','review','delivered']
@app.post('/api/projects/{pid}/meta')
def meta(pid:str,body:dict):
    pr=_load(pid)
    for k in ('name','client','measurer','reviewer'):
        if k in body: pr[k]=str(body[k]).strip()[:80]
    if 'due' in body: pr['due']=str(body['due'])[:10]
    if 'checks' in body: pr['checks'].update({k:bool(v) for k,v in body['checks'].items()})
    if 'status' in body:
        if body['status'] not in STATUSES: raise HTTPException(400,'סטטוס לא מוכר')
        if body['status']=='review' and not pr.get('boq'): raise HTTPException(400,'אי אפשר לשלוח לבדיקה לפני שחושב כתב')
        if body['status']=='delivered':
            if not all(c['ok'] for c in _checks(pr)): raise HTTPException(400,'לא כל הבדיקות הושלמו')
            pr['delivered_at']=datetime.datetime.now().isoformat(timespec='minutes')
        if pr['status']!=body['status']: _ev(pr,{'received':'הוחזר לשלב קבלה','measuring':'הוחזר למדידה','review':'נשלח לבדיקה','delivered':'נמסר ללקוח'}[body['status']])
        pr['status']=body['status']
    _save(pr); return pr

def _checks(pr):
    ck=pr.get('checks',{}); sh=pr['sheets']; needs=pr.get('needs',[])
    return [dict(k='scale',t='קנה מידה כויל מהתוכנית',auto=True,ok=bool(pr.get('scale')) and not any(n['t'].startswith('קנה מידה') for n in needs)),
     dict(k='sheets',t='כל הגיליונות סווגו',auto=True,ok=bool(sh) and all(x['type']!='unknown' for x in sh)),
     dict(k='high',t='אין ממצאים דחופים פתוחים',auto=True,ok=not any(n['p']=='high' for n in needs),note=f"{sum(1 for n in needs if n['p']=='high')} פתוחים"),
     dict(k='asm',t='ההנחות אושרו',auto=False,ok=bool(ck.get('asm'))),
     dict(k='second',t='בדיקה שנייה בוצעה',auto=False,ok=bool(ck.get('second')),note=pr.get('reviewer') or 'לא הוגדר בודק')]

@app.get('/api/projects/{pid}')
def get(pid:str): return _load(pid)

@app.delete('/api/projects/{pid}')
def delete(pid:str): shutil.rmtree(_pdir(pid),ignore_errors=True); return dict(ok=True)

@app.post('/api/projects/{pid}/upload')
async def upload(pid:str,files:list[UploadFile]=File(...)):
    pr=_load(pid); rules=_rules(); seen={s['hash']:s for s in pr['sheets']}; added=[]
    for f in files:
        if not f.filename.lower().endswith('.pdf'): added.append(dict(file=f.filename,error='רק קבצי PDF נתמכים כרגע (DWG בהמשך)')); continue
        sid=uuid.uuid4().hex[:6]; path=os.path.join(_pdir(pid),'files',sid+'.pdf')
        with open(path,'wb') as o: o.write(await f.read())
        h=C.file_hash(path)
        if h in seen:
            os.remove(path); added.append(dict(file=f.filename,dup_of=seen[h]['file'])); continue
        try: ft=C.features(path); cl=C.classify(ft,rules)
        except Exception as e: os.remove(path); added.append(dict(file=f.filename,error='הקובץ לא נקרא: '+str(e)[:80])); continue
        s=dict(sid=sid,file=f.filename,hash=h,w_mm=ft['w_mm'],h_mm=ft['h_mm'],n_draw=ft['n_draw'],layers=ft['layers'],type=cl['type'],conf=cl['conf'],why=cl['why'],floor=None,confirmed=False,version=1)
        try: rv=C.revision_of(ft,path,pr['sheets'],_pdir(pid))
        except Exception: rv=None
        if rv:
            s['pending_of']=rv[0]['sid']; s['sim']=round(rv[1]*100); pr['pending']=[x for x in pr['pending'] if x['hash']!=h]+[s]; seen[h]=s
            added.append(dict(file=f.filename,sid=sid,type=s['type'],conf=s['conf'],revision_of=rv[0]['sid'],sim=s['sim'])); continue
        pr['sheets'].append(s); seen[h]=s; added.append(dict(file=f.filename,sid=sid,type=s['type'],conf=s['conf'],disc=C.DISC.get(s['type'],'')))
    n_ok=sum(1 for a in added if a.get('sid') and not a.get('revision_of'))
    if n_ok: _stash(pr); _ev(pr,f'הועלו {n_ok} תוכניות')
    if n_ok: pr['stale']=True
    _save(pr); return dict(added=added,project=pr)

@app.post('/api/projects/{pid}/sheets/{sid}')
def label(pid:str,sid:str,body:dict):
    pr=_load(pid); s=next((x for x in pr['sheets'] if x['sid']==sid),None)
    if not s: raise HTTPException(404,'גיליון לא נמצא')
    if 'type' in body:
        if body['type'] not in C.TYPES: raise HTTPException(400,'סוג לא מוכר')
        s['type']=body['type']; s['confirmed']=True; s['conf']='confirmed'; s['why']='אושר על ידך'
        rules=[r for r in _rules() if set(r['sig'])!=set(C.sig(s['layers']))]; rules.append(dict(sig=C.sig(s['layers']),type=s['type']))
        os.makedirs(ROOT,exist_ok=True); json.dump(rules,open(RULES,'w',encoding='utf-8'),ensure_ascii=False)
    if 'type' in body and 'disc' not in body: s.pop('disc',None)
    if 'disc' in body:
        if body['disc'] and body['disc'] not in dict(C.DISCS): raise HTTPException(400,'תחום לא מוכר')
        s['disc']=body['disc'] or None
    if 'floor' in body: s['floor']=body['floor'] or None
    _stash(pr); pr['stale']=True; _save(pr); return pr

def _snap(rows): return {r['desc']:dict(q=r.get('q'),unit=r.get('unit')) for r in (rows or []) if r.get('kind')=='item' and r.get('q') is not None}
def _stash(pr):
    if pr.get('boq') and pr.get('prev_snap') is None: pr['prev_snap']=_snap(pr['boq']); pr['prev_ver']=pr.get('version',1)

@app.post('/api/projects/{pid}/revision')
def revision(pid:str,body:dict):
    """decide what to do with an uploaded sheet that looks like a new revision: replace | add | skip"""
    pr=_load(pid); n=next((x for x in pr['pending'] if x['sid']==body.get('sid')),None)
    if not n: raise HTTPException(404,'אין גיליון ממתין')
    act=body.get('action'); pr['pending']=[x for x in pr['pending'] if x['sid']!=n['sid']]
    old=next((x for x in pr['sheets'] if x['sid']==n.get('pending_of')),None)
    n.pop('pending_of',None); n.pop('sim',None)
    if act=='skip':
        try: os.remove(os.path.join(_pdir(pid),'files',n['sid']+'.pdf'))
        except OSError: pass
        _save(pr); return pr
    _stash(pr)
    if act=='replace' and old:
        for k in ('type','floor','disc','confirmed','conf','why'):
            if k in old: n[k]=old[k]
        n['version']=old.get('version',1)+1; n['prev']=[dict(sid=old['sid'],file=old['file'],version=old.get('version',1))]+old.get('prev',[])
        pr['sheets']=[n if x['sid']==old['sid'] else x for x in pr['sheets']]; _ev(pr,f"גיליון עודכן לגרסה {n['version']}: {n['file'][:30]}")
    elif act in ('add','replace'): pr['sheets'].append(n); _ev(pr,'נוסף גיליון חדש')
    else: raise HTTPException(400,'פעולה לא מוכרת')
    pr['stale']=True; _save(pr); return pr

@app.get('/api/projects/{pid}/sheets/{sid}/view.png')
def view(pid:str,sid:str):
    import fitz
    if not re.fullmatch(r'[a-f0-9]{6}',sid): raise HTTPException(404,'')
    src=os.path.join(_pdir(pid),'files',sid+'.pdf'); out=os.path.join(_pdir(pid),'files',sid+'_big.png')
    if not os.path.exists(src): raise HTTPException(404,'')
    if not os.path.exists(out):
        pg=fitz.open(src)[0]; z=min(3000/pg.rect.width,4.0); pg.get_pixmap(matrix=fitz.Matrix(z,z),alpha=False).save(out)
    return FileResponse(out,media_type='image/png')

FLOOR={'ground':'קומת קרקע','first':'קומה א׳','second':'קומה ב׳','roof':'גג'}

@app.post('/api/projects/{pid}/run')
def run(pid:str):
    pr=_load(pid); sheets=pr['sheets']; log=[]; lines=[]; assumptions=[]
    path=lambda s:os.path.join(_pdir(pid),'files',s['sid']+'.pdf')
    scale=None
    for s in sheets:
        if s['type']=='foundation_piles':
            sc,ok,tot=E.calibrate_scale(path(s))
            if sc: scale=sc; log.append(f"כיול קנה מידה מכלונסאות: 1:{sc} ({ok} מתוך {tot} עיגולים בקטרים תקניים)"); break
    if not scale:
        scale=50; assumptions.append(dict(t='קנה מידה 1:50 הונח',why='לא נמצאה תוכנית כלונסאות לכיול. כל המדידות נשענות על ההנחה הזו',how='אשר קנה מידה, או העלה תוכנית יסודות'))
    pr['scale']=scale
    fps=[s for s in sheets if s['type']=='floor_plan']
    slab_areas=[]; labels=_jl(SYMF,{}); lmap=_jl(LMAPF,{})
    for s in fps:
        if not s['floor']: assumptions.append(dict(t=f"באיזו קומה גיליון \"{s['file']}\"?",why='אי אפשר לקרוא את שם הקומה מהגיליון (טקסט וקטורי)',gives='תיאור הסעיפים ואיטום גג',how='בחר קומה בטבלת הגיליונות'))
    for s in sheets:
        t=s['type']
        try:
            if t=='foundation_piles':
                l,info=E.piles(path(s),scale); lines+=l; log.append(f"{s['file']}: כלונסאות {info}")
            elif t=='floor_plan':
                fl=FLOOR.get(s['floor'],'קומה לא מסומנת')
                if s['floor']=='roof':
                    l,info=E.roof(path(s),scale); lines+=l; log.append(f"{s['file']}: גג {info}")
                    if info.get('area_m2'): slab_areas.append(info['area_m2'])
                    l,info=E.walls(path(s),scale,'מעקות גג'); lines+=l
                else:
                    l,info=E.walls(path(s),scale,fl); lines+=l; log.append(f"{s['file']}: קירות {info}")
                    try:
                        _,ri=E.roof(path(s),scale)
                        if ri.get('area_m2'): slab_areas.append(ri['area_m2'])
                    except Exception: pass
                    if info.get('extra_plan'): assumptions.append(dict(t=f"בגיליון \"{s['file']}\" זוהו שתי תוכניות דומות",why='נבחרה הגדולה מביניהן',how='אשר שזו הנכונה'))
            elif t=='elevation':
                l,info=E.openings(path(s),scale); lines+=l; log.append(f"{s['file']}: פתחים {info}")
            elif t=='electrical':
                l,info=E.rooms(path(s),scale)
                nel=sum(1 for t in sheets[:sheets.index(s)] if t['type']=='electrical')+1
                for x in l: x['desc']+=f" (תוכנית חשמל {nel})"
                lines+=l; log.append(f"{s['file']}: חללים {info}")
                if labels:
                    sl,si=SY.count(path(s),labels); lines+=sl; log.append(f"{s['file']}: סמלי חשמל {si}")
        except Exception as e: log.append(f"{s['file']}: שגיאה בחישוב: {str(e)[:100]}")
        for layer,m in (lmap.get(t) or {}).items():
            try:
                q=MP.measure(path(s),scale,layer,m['mode'])
                if q: lines.append(dict(ch=int(m['ch']),sub=int(m.get('sub',1)),desc=m['desc'],unit=m['unit'],q=q,st='נמדד',src=f"מיפוי שכבה {layer} ({s['file'][:18]})",need='',cover=[m['desc']]))
            except Exception as e: log.append(f"{s['file']}: שגיאת מיפוי שכבה {layer}: {str(e)[:60]}")
    P=BR.params(pr)
    lines+=DV.derive(lines,slab_areas,P)
    merged={}
    for l in lines:
        k=(l['ch'],l.get('sub',1),l['desc'],l['unit'])
        if k in merged and l['q'] is not None and merged[k]['q'] is not None: merged[k]['q']=round(merged[k]['q']+l['q'],2)
        else: merged[k]=l
    lines=list(merged.values())
    pr['lines']=lines; pr['asm']=assumptions; pr['log']=log
    rebuild(pr)
    if pr.get('prev_snap') is not None:
        old=pr['prev_snap']; new=_snap(pr['boq']); ch=[]
        for d in sorted(set(old)|set(new)):
            a=old.get(d,{}).get('q'); b=new.get(d,{}).get('q')
            if a!=b: ch.append(dict(desc=d,unit=(new.get(d) or old.get(d)).get('unit'),old=a,new=b))
        pr['version']=pr.get('prev_ver',1)+1
        pr['changes']=dict(frm=pr['version']-1,to=pr['version'],at=datetime.datetime.now().isoformat(timespec='minutes'),items=ch)
        pr['prev_snap']=None; _ev(pr,f"גרסה {pr['version']}: {len(ch)} סעיפים השתנו")
    else: _ev(pr,'כתב הכמויות חושב')
    if pr.get('status')=='received': pr['status']='measuring'
    pr['stale']=False; _save(pr)
    return pr

def rebuild(pr):
    """BOQ = engine lines + manual items, with quantity overrides and hidden items; then materials and needs"""
    P=BR.params(pr); sheets=pr['sheets']; lines=list(pr.get('lines',[]))
    for m in pr.get('manual',[]): lines.append(dict(ch=m['ch'],sub=m.get('sub',1),desc=m['desc'],unit=m['unit'],q=m['q'],st='ידני',src='הוזן ידנית',need=''))
    present={s['type'] for s in sheets}; needed={e[4] for e in B.EXPECTED}
    missing=sorted(t for t in needed if t not in present)
    unbuilt=sorted(t for t in present if t not in B.BUILT and t in needed or t=='unknown' and t in present)
    rows=B.build(pr['name'],lines,present,pr.get('ov'),set(pr.get('hidden',[])),pr.get('chnames'))
    mats,used=BR.explode(rows,P); asm=list(pr.get('asm',[]))
    for k in used: asm.append(dict(t=f"הנחה: {BR.DEFAULT_PARAMS[k]['label']} = {P[k]}",why=BR.DEFAULT_PARAMS[k]['why'],gives='פירוק לחומרים',how='עדכן בהגדרות הפרויקט'))
    pr['boq']=rows; pr['materials']=mats; pr['needs']=B.needs(sheets,missing,unbuilt,asm)

@app.post('/api/projects/{pid}/builder')
def builder(pid:str,body:dict):
    """add / quantity override / hide / restore / rename chapter, without re-running the engines"""
    pr=_load(pid); a=body.get('action')
    if 'lines' not in pr: pr['lines']=[]
    if a=='add':
        if not body.get('desc') or not body.get('unit'): raise HTTPException(400,'חסר תיאור או יחידה')
        pr.setdefault('manual',[]).append(dict(ch=int(body.get('ch') or 99),sub=int(body.get('sub') or 1),desc=body['desc'].strip(),unit=body['unit'].strip(),q=float(body['q']) if body.get('q') not in ('',None) else None))
        if body.get('chname'): pr.setdefault('chnames',{})[str(int(body['ch']))]=body['chname']
    elif a=='qty':
        pr.setdefault('ov',{})[body['desc']]=float(body['q']) if body.get('q') not in ('',None) else None
    elif a=='hide': pr.setdefault('hidden',[]).append(body['desc'])
    elif a=='unhide': pr['hidden']=[d for d in pr.get('hidden',[]) if d!=body['desc']]
    elif a=='rm_manual': pr['manual']=[m for m in pr.get('manual',[]) if m['desc']!=body['desc']]
    else: raise HTTPException(400,'פעולה לא מוכרת')
    rebuild(pr); _save(pr); return pr

@app.post('/api/projects/{pid}/start')
def start_empty(pid:str):
    """start a BOQ without plans: only the expected items, to be filled by hand"""
    pr=_load(pid); pr['lines']=pr.get('lines',[]); pr['asm']=[]; rebuild(pr); _save(pr); return pr

@app.get('/api/projects/{pid}/sheets/{sid}/thumb.png')
def thumb(pid:str,sid:str):
    import fitz
    if not re.fullmatch(r'[a-f0-9]{6}',sid): raise HTTPException(404,'')
    src=os.path.join(_pdir(pid),'files',sid+'.pdf'); out=os.path.join(_pdir(pid),'files',sid+'.png')
    if not os.path.exists(src): raise HTTPException(404,'')
    if not os.path.exists(out):
        pg=fitz.open(src)[0]; z=min(900/pg.rect.width,2.0); pg.get_pixmap(matrix=fitz.Matrix(z,z),alpha=False).save(out)
    return FileResponse(out,media_type='image/png')

@app.get('/api/projects/{pid}/export.xlsx')
def export(pid:str):
    pr=_load(pid)
    if not pr.get('boq'): raise HTTPException(400,'עוד לא חושב כתב')
    out=os.path.join(_pdir(pid),'boq.xlsx'); B.to_xlsx(pr['boq'],pr['name'],out,pr['needs'],prices=pr['prices'],meta=dict(client=pr.get('client',''),version='v%s'%pr.get('version',1),measurer=pr.get('measurer',''),reviewer=pr.get('reviewer',''),status=pr.get('status',''))); return FileResponse(out,filename='kamuyot_boq.xlsx')



# ---------- pricing, memory, bills, orders ----------
MEM=os.path.join(ROOT,'prices_memory.json')
def _mem():
    try: return json.load(open(MEM,encoding='utf-8'))
    except Exception: return []

@app.post('/api/projects/{pid}/params')
def set_params(pid:str,body:dict):
    pr=_load(pid); pr['params']={k:float(v) for k,v in body.items() if k in BR.DEFAULT_PARAMS}; _save(pr); return run(pid) if pr.get('boq') else pr

from .spec import SPEC, CHAPTERS
@app.get('/api/spec')
def spec(): return dict(chapters={str(k):v for k,v in CHAPTERS.items()},items=SPEC)

ENGINES=[
 dict(k='piles',name='כלונסאות',ch='23',st='active',does='סופר כלונסאות לפי קוטר מתוכנית היסודות ומכייל את קנה המידה',needs='אורך כלונס (OCR של טבלת מפלסים) כדי לחשב בטון וזיון מדויקים'),
 dict(k='walls',name='קירות',ch='04',st='partial',does='מודד אורך קירות לפי עובי מתוכנית קומה, לפי קומה',needs='הפרדה בין בלוק לבטון, וגובה קומה מהחתכים'),
 dict(k='rooms',name='חללים וריצוף',ch='10',st='active',does='מזהה חדרים סגורים ומחשב שטח נטו',needs='פיצול חללים פתוחים ושמות חדרים'),
 dict(k='roof',name='גג ואיטום',ch='05',st='active',does='שטח גג לפי קו המעקה החיצוני',needs='-'),
 dict(k='openings',name='פתחים',ch='06',st='partial',does='סופר פתחים ושטחם מחזיתות',needs='פרט דלתות וחלונות לפי תגיות (OCR)'),
 dict(k='plaster',name='טיח וצבע',ch='09, 11',st='estimate',does='הערכה מאורך קירות כפול גובה פחות פתחים',needs='היקף חדרים מדויק ושטחי תקרה'),
 dict(k='concrete',name='בטון, תבניות וזיון',ch='02',st='estimate',does='נגזר: קירות בטון, תקרות לפי שטח ועובי שהונח, כלונסאות לפי אורך שהונח, ותבניות וזיון משקל לפי ק"ג למ"ק',needs='מידות אמיתיות וטבלאות זיון מגיליונות הקונסטרוקציה (OCR)'),
 dict(k='symbols',name='ספריית סמלי חשמל',ch='08',st='partial',does='מקבץ סמלים לפי צורה. מסמנים כל צורה פעם אחת (תאורה, שקע, מפסק...) והיא נספרת בכל תוכנית חשמל',needs='סימון הצורות על ידך. כל תיוג נשמר ומשפר את הדיוק'),
 dict(k='mapper',name='מיפוי שכבות כללי',ch='כל הפרקים',st='active',does='כל שכבה בכל תוכנית PDF מדידה לפי אורך, שטח או ספירה ומשויכת לסעיף. מתאים לאינסטלציה, מיזוג, פיתוח וכל תוכנית עם שכבות',needs='לבחור שכבה וסעיף פעם אחת. ההוראה נשמרת לפי סוג גיליון'),
 dict(k='plumbing',name='אינסטלציה, מיזוג ופיתוח',ch='07, 15, 40, 57',st='partial',does='עובד דרך מיפוי שכבות ברגע שיש תוכנית מהסוג הזה',needs='תוכניות לדוגמה להכרת השכבות ולאימון זיהוי אוטומטי'),
]
@app.get('/api/engines')
def engines_list(): return ENGINES

@app.get('/legacy')
def legacy():
    p=os.environ.get('KAMUYOT_LEGACY') or os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),'kamuyot_system_empty.html')
    if not os.path.exists(p): raise HTTPException(404,'אין קובץ')
    return FileResponse(p,media_type='text/html')

@app.get('/api/projects/{pid}/sheets/{sid}/symbols')
def sheet_symbols(pid:str,sid:str):
    if not re.fullmatch(r'[a-f0-9]{6}',sid): raise HTTPException(404,'')
    p=os.path.join(_pdir(pid),'files',sid+'.pdf')
    if not os.path.exists(p): raise HTTPException(404,'')
    labels=_jl(SYMF,{}); g=SY.groups(p)
    return dict(labels=SY.LABEL_NAMES,groups=[dict(g=x,label=labels.get(x['sig'])) for x in g if x['n']>=2][:40],total=sum(x['n'] for x in g),labelled=sum(x['n'] for x in g if labels.get(x['sig'])))

@app.get('/api/projects/{pid}/sheets/{sid}/crop.png')
def sheet_crop(pid:str,sid:str,x:float,y:float,w:float,h:float):
    if not re.fullmatch(r'[a-f0-9]{6}',sid): raise HTTPException(404,'')
    p=os.path.join(_pdir(pid),'files',sid+'.pdf'); out=os.path.join(_pdir(pid),'files',f'{sid}_{int(x)}_{int(y)}.png')
    if not os.path.exists(p): raise HTTPException(404,'')
    if not os.path.exists(out): SY.crop_png(p,x,y,w,h,out)
    return FileResponse(out,media_type='image/png')

@app.post('/api/symbols')
def label_symbol(body:dict):
    lab=_jl(SYMF,{})
    if body.get('label') and body['label'] in SY.LABEL_NAMES: lab[body['sig']]=body['label']
    else: lab.pop(body.get('sig'),None)
    _js(SYMF,lab); return dict(ok=True,n=len(lab))

@app.get('/api/projects/{pid}/sheets/{sid}/layers')
def sheet_layers(pid:str,sid:str):
    pr=_load(pid); s=next((x for x in pr['sheets'] if x['sid']==sid),None)
    if not s: raise HTTPException(404,'')
    cache=os.path.join(_pdir(pid),'files',sid+'.layers.json'); sc=pr.get('scale') or 50
    st=_jl(cache,None)
    if not st or st.get('scale')!=sc:
        st=dict(scale=sc,layers=MP.stats(os.path.join(_pdir(pid),'files',sid+'.pdf'),sc)); _js(cache,st)
    return dict(scale=sc,type=s['type'],layers=st['layers'],map=_jl(LMAPF,{}).get(s['type'],{}))

@app.post('/api/layermap')
def layermap(body:dict):
    m=_jl(LMAPF,{}); t=m.setdefault(body['type'],{})
    if body.get('remove'): t.pop(body['layer'],None)
    else:
        if not body.get('desc') or body.get('mode') not in ('length','area','count','circles'): raise HTTPException(400,'חסר תיאור או שיטת מדידה')
        t[body['layer']]=dict(desc=body['desc'].strip(),ch=int(body.get('ch') or 99),sub=int(body.get('sub') or 1),unit=body.get('unit') or 'מטר',mode=body['mode'])
    _js(LMAPF,m); return dict(ok=True)

@app.get('/api/params')
def params_def(): return BR.DEFAULT_PARAMS

@app.post('/api/projects/{pid}/prices')
def set_prices(pid:str,body:dict):
    pr=_load(pid); pr.setdefault('prices',{}).update({k:float(v) for k,v in (body.get('items') or {}).items() if v not in ('',None)})
    pr.setdefault('mprices',{}).update({k:float(v) for k,v in (body.get('materials') or {}).items() if v not in ('',None)})
    if body.get('remember'):
        mem=_mem(); idx={(m['desc'],m['unit']):m for m in mem}
        for r in pr['boq'] or []:
            if r['kind']=='item' and r['code'] in pr['prices']: idx[(r['desc'],r['unit'])]=dict(desc=r['desc'],unit=r['unit'],price=pr['prices'][r['code']])
        json.dump(list(idx.values()),open(MEM,'w',encoding='utf-8'),ensure_ascii=False)
    _save(pr); return pr

@app.post('/api/projects/{pid}/suggest')
def suggest(pid:str):
    pr=_load(pid); mem=_mem(); out={}
    for r in pr['boq'] or []:
        if r['kind']=='item' and r['code'] not in pr['prices']:
            s=BR.suggest(r['desc'],r['unit'],mem)
            if s: out[r['code']]=s
    return out

@app.post('/api/prices/import')
async def import_prices(file:UploadFile=File(...)):
    """learn prices from any BOQ xlsx: rows with description, unit and price"""
    import openpyxl, io
    wb=openpyxl.load_workbook(io.BytesIO(await file.read()),data_only=True); mem={(m['desc'],m['unit']):m for m in _mem()}; n=0
    for ws in wb:
        for row in ws.iter_rows(values_only=True):
            c=[x for x in row]
            for i in range(len(c)-3):
                d,u,q,p=c[i+1] if i+1<len(c) else None,c[i+2] if i+2<len(c) else None,c[i+3] if i+3<len(c) else None,c[i+4] if i+4<len(c) else None
                if isinstance(d,str) and len(d)>8 and isinstance(u,str) and len(u.strip())<8 and isinstance(p,(int,float)) and p>0:
                    mem[(d.strip(),u.strip())]=dict(desc=d.strip(),unit=u.strip(),price=float(p)); n+=1; break
    json.dump(list(mem.values()),open(MEM,'w',encoding='utf-8'),ensure_ascii=False); return dict(learned=n,total=len(mem))

@app.get('/api/projects/{pid}/summary')
def summary(pid:str):
    pr=_load(pid); rows=pr.get('boq') or []; items=[r for r in rows if r['kind']=='item']
    P=BR.params(pr); tot=BR.totals(rows,pr['prices'])
    by={}
    for r in rows:
        if r['kind']=='ch': cur=r['desc']; by[cur]=dict(total=0,items=0,measured=0,priced=0)
        elif r['kind']=='item' and rows.index(r)>0:
            b=by[cur]; b['items']+=1; b['measured']+= 1 if r.get('q') else 0
            if r.get('q') and pr['prices'].get(r['code']): b['priced']+=1; b['total']+=r['q']*pr['prices'][r['code']]
    by.pop(pr['name'],None)
    billed=sum(b['total'] for b in pr['bills']) if pr['bills'] else 0
    return dict(name=pr['name'],sheets=len(pr['sheets']),items=len(items),measured=sum(1 for r in items if r.get('q')),priced=sum(1 for r in items if r.get('q') and pr['prices'].get(r['code'])),total=tot,chapters=by,bills=len(pr['bills']),billed=billed,orders=len(pr['orders']),needs=len(pr['needs']),
      id=pr['id'],client=pr['client'],status=pr['status'],measurer=pr['measurer'],reviewer=pr['reviewer'],due=pr['due'],version=pr['version'],updated=pr['updated'],created=pr['created'],delivered_at=pr['delivered_at'],
      trust=dict(measured=sum(1 for r in items if r.get('q') is not None and r.get('st')=='נמדד'),manual=sum(1 for r in items if r.get('q') is not None and r.get('st')=='ידני'),est=sum(1 for r in items if r.get('q') is not None and r.get('st') not in ('נמדד','ידני')),missing=sum(1 for r in items if r.get('q') is None)),
      thumb_sid=next((x['sid'] for t in ('floor_plan','foundation_piles') for x in pr['sheets'] if x['type']==t),(pr['sheets'][0]['sid'] if pr['sheets'] else None)),checks=_checks(pr),changes=pr.get('changes'),events=pr.get('events',[])[-8:],high=sum(1 for n in pr['needs'] if n['p']=='high'))

@app.post('/api/projects/{pid}/bills')
def new_bill(pid:str,body:dict):
    """body: {pct:{code:cumulative_pct}}"""
    pr=_load(pid); P=BR.params(pr); prev=pr['bills'][-1]['pct'] if pr['bills'] else {}
    cum={k:float(v) for k,v in body.get('pct',{}).items()}
    for k,v in prev.items(): cum.setdefault(k,v)
    b=BR.bill(pr['boq'],pr['prices'],cum,prev,P); b.update(n=len(pr['bills'])+1,date=str(datetime.date.today()),pct=cum,total_this=b['this'])
    pr['bills'].append(dict(n=b['n'],date=b['date'],pct=cum,total=b['total'],this=b['this'],retention=b['retention'],vat=b['vat'],net=b['net'])); _save(pr); return b

@app.get('/api/projects/{pid}/bills/{n}')
def get_bill(pid:str,n:int):
    pr=_load(pid); P=BR.params(pr); i=n-1
    if i<0 or i>=len(pr['bills']): raise HTTPException(404,'חשבון לא נמצא')
    prev=pr['bills'][i-1]['pct'] if i else {}; b=BR.bill(pr['boq'],pr['prices'],pr['bills'][i]['pct'],prev,P); b.update(n=n,date=pr['bills'][i]['date']); return b

@app.post('/api/projects/{pid}/orders')
def new_order(pid:str,body:dict):
    pr=_load(pid)
    items=[dict(material=i['material'],unit=i['unit'],qty=float(i['qty']),price=pr['mprices'].get(i['material'])) for i in body.get('items',[]) if float(i.get('qty',0))>0]
    if not items: raise HTTPException(400,'לא נבחרו חומרים')
    o=dict(n=len(pr['orders'])+1,supplier=(body.get('supplier') or 'ספק לא צוין').strip(),needed=body.get('needed',''),note=body.get('note',''),date=str(datetime.date.today()),status='טיוטה',items=items)
    pr['orders'].append(o); _save(pr); return o

@app.post('/api/projects/{pid}/orders/{n}')
def order_status(pid:str,n:int,body:dict):
    pr=_load(pid); o=next((x for x in pr['orders'] if x['n']==n),None)
    if not o: raise HTTPException(404,'הזמנה לא נמצאה')
    if body.get('status') in ('טיוטה','נשלחה','סופקה חלקית','סופקה','בוטלה'): o['status']=body['status']
    _save(pr); return o

@app.get('/api/projects/{pid}/quote.xlsx')
def quote(pid:str):
    pr=_load(pid)
    if not pr.get('boq'): raise HTTPException(400,'עוד לא חושב כתב')
    rows=[dict(r) for r in pr['boq']]
    import copy
    out=os.path.join(_pdir(pid),'quote.xlsx'); B.to_xlsx(rows,pr['name'],out,pr['needs'],prices=pr['prices'],meta=dict(client=pr.get('client',''),version='v%s'%pr.get('version',1),measurer=pr.get('measurer',''),reviewer=pr.get('reviewer',''),status=pr.get('status',''))); return FileResponse(out,filename='kamuyot_quote.xlsx')

@app.get('/api/projects/{pid}/orders/{n}.xlsx')
def order_xlsx(pid:str,n:int):
    import openpyxl
    from openpyxl.styles import Font
    pr=_load(pid); o=next((x for x in pr['orders'] if x['n']==n),None)
    if not o: raise HTTPException(404,'הזמנה לא נמצאה')
    wb=openpyxl.Workbook(); ws=wb.active; ws.sheet_view.rightToLeft=True; ws.title='הזמנה'
    ws.append([f"הזמנת רכש {n} - {pr['name']}"]); ws['A1'].font=Font(bold=True,size=16); ws.append(['ספק',o['supplier']]); ws.append(['תאריך אספקה נדרש',o['needed']]); ws.append(['הערה',o['note']]); ws.append([])
    ws.append(['חומר',"יח'",'כמות','מחיר','סה"כ'])
    for c in ws[6]: c.font=Font(bold=True)
    for i in o['items']:
        r=ws.max_row+1; ws.append([i['material'],i['unit'],i['qty'],i['price'] or 0,f'=C{r}*D{r}'])
    ws.append(['','','','סה"כ',f'=SUM(E7:E{ws.max_row})'])
    for col,w in zip('ABCDE',[40,10,12,12,14]): ws.column_dimensions[col].width=w
    out=os.path.join(_pdir(pid),f'order{n}.xlsx'); wb.save(out); return FileResponse(out,filename=f'order{n}.xlsx')

STATIC=os.path.join(os.path.dirname(os.path.abspath(__file__)),'static')
if os.path.isdir(STATIC): app.mount('/',StaticFiles(directory=STATIC,html=True),name='ui')
