import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
wb=openpyxl.Workbook()
H=PatternFill('solid',fgColor='1F3A5F'); HF=Font(bold=True,color='FFFFFF',name='Arial')
ST={'נמדד':'D8F0DC','חלקי':'FFF1CC','דורש השלמה':'F8D7D7'}
thin=Side(style='thin',color='BBBBBB')
def sheet(ws,title,cols,rows,widths):
    ws.title=title; ws.sheet_view.rightToLeft=True
    ws.append(cols)
    for c in ws[1]: c.fill=H; c.font=HF; c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    for r in rows: ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment=Alignment(vertical='top',wrap_text=True); c.border=Border(top=thin,bottom=thin,left=thin,right=thin); c.font=Font(name='Arial',size=10)
        st=row[-3].value if title!='סיכום' else None
    ws.freeze_panes='A2'
    for i,w in enumerate(widths): ws.column_dimensions[chr(65+i)].width=w
# (פרק, סעיף, כמות, יחידה, סטטוס, מקור/שיטה, הערה)
exec(open(__import__('os').path.dirname(__import__('os').path.abspath(__file__))+'/villa_data.py',encoding='utf-8').read())
ws=wb.active
sheet(ws,'כמויות',['פרק','סעיף','כמות','יחידה','סטטוס','מקור ושיטה','הערה'],Q,[8,48,12,12,14,52,52])
for row in ws.iter_rows(min_row=2):
    c=row[4]; c.fill=PatternFill('solid',fgColor=ST[c.value]); c.font=Font(bold=True,name='Arial',size=10)
S=wb.create_sheet()
sheet(S,'גיליונות',['קובץ','סוג','מה בו','קריאות'],[
['A4','4 חזיתות','פתחים בחזית צפונית','וקטורי'],['A5','3 חתכים','גבהי קומות','וקטורי'],['A6','תוכנית פיתוח','גדרות ושער','וקטורי'],['A1–A3','אדריכלות/קונסטרוקציה כולל','קומת קרקע, קומה א׳, גג · 1:50','וקטורי · שכבות · טקסט כקווים'],
['A7, A8','פרטי בנייה','איטום וחיבורי מעטפת','וקטורי'],
['A12, A13','פרטי דלתות וחלונות','24 פרטים','וקטורי'],
['kon7-1','יסודות','27 כלונסאות','נמדד'],
['kon7-2…8','קונסטרוקציה','תקרות, קורות, פרטים, מדרגות','טרם מדוד'],
['חשמל ×4','חשמל','קרקע, גג · 2 עותקים כפולים','טקסט אמיתי'],
],[16,26,40,28])
wb.save('villa_ivgi_quantities.xlsx')
